"""
第3课实战：多格式文档解析（2026-10 现代化版本）

在第2课基础上迭代（对照课程 rag_app_lesson3.py）：
1. 新增多格式解析：支持 PDF/TXT/DOC/DOCX/PPTX/XLSX/CSV/MD/XML/HTML
2. 索引流程从"单文件"升级为"遍历整个文件夹"

与2024课程示例的关键差异：
- 课程用 langchain_community 的 10 种 Document Loader（PDFPlumberLoader等），
  该包已于2026年6月被官方弃用归档，且Loader无官方替代包。
  我们的方案：保持课程"扩展名→解析器"分发表的思想不变，
  每种格式直接调用底层解析库（它们正是各种Loader内部封装的库）：
    .pdf  → pdfplumber   （课程 PDFPlumberLoader 的底层库，中文支持好、表格解析强）
    .txt/.md → 直接读取  （本来就是纯文本）
    .doc  → macOS自带textutil命令（旧版二进制Word格式；Linux上需LibreOffice）
    .docx → python-docx  （课程 UnstructuredWordDocumentLoader 的底层方案）
    .pptx → python-pptx
    .xlsx → openpyxl
    .csv  → 内置csv模块
    .html → BeautifulSoup
    .xml  → 内置xml.etree
- 检索/生成流程与第2课完全一致（课程同样如此迭代）

运行方式：
    uv run --env-file .env python lesson3_multiformat_rag.py
"""

import os

# 必须在导入 sentence-transformers 之前设置：
# 禁用分词器并行（多线程/多进程场景下会死锁或告警），课程代码同样处理
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")

import csv
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

import docx  # python-docx 包的导入名是 docx
import faiss
import numpy as np
import openpyxl
import pdfplumber
import pptx  # python-pptx 包的导入名是 pptx
from bs4 import BeautifulSoup
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sentence_transformers import SentenceTransformer

# ---------- 配置区 ----------
# 用"当前文件所在目录"拼出绝对路径，从任何目录启动都能找到文件
BASE_DIR = Path(__file__).resolve().parent
EMBEDDING_MODEL_PATH = BASE_DIR / "models" / "bge-small-zh-v1.5"
# 第3课的测试数据文件夹：10种格式的文件内容完全相同，用于对比解析效果
DATA_FOLDER = BASE_DIR / "data" / "lesson3"

CHUNK_SIZE = 512  # 每个文本块的最大字符数（受Embedding模型512 token输入上限约束）
CHUNK_OVERLAP = 128  # 相邻文本块之间的重叠字符数
TOP_K = 3  # 检索返回最相似的前K个文本块


# ======================================================================
# 第一部分：每种文档格式的解析函数
# 课程把这些逻辑封装在 langchain_community 的 Loader 类里；
# 我们直接调用底层解析库，每种格式一个小函数，内部逻辑完全透明
# ======================================================================


def load_pdf(file_path):
    """
    解析PDF文件：pdfplumber 逐页提取文本
    （pdfplumber 对中文支持好、表格解析强，但对双栏排版解析较弱；
    更强的PDF解析方案——版面分析/扫描件OCR——见本课讲解和第4课）
    """
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            # 个别页面可能是纯图片页，extract_text会返回None，需要跳过
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def load_txt(file_path):
    """解析TXT文件：纯文本直接读取"""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def load_md(file_path):
    """
    解析Markdown文件：同样是纯文本直接读取
    （Markdown正是文档解析领域公认的"统一输出格式"，无需再转换）
    """
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def load_doc(file_path):
    """
    解析旧版Word(.doc)文件：二进制格式，python-docx不支持。
    macOS 自带 textutil 命令行工具可以直接转换（无需安装700MB的LibreOffice；
    Linux 服务器上部署时则用 LibreOffice 的 soffice 命令替代）
    """
    try:
        # textutil -convert txt -stdout：转换为纯文本并打印到标准输出
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", file_path],
            capture_output=True,  # 捕获输出而不是直接打印
            text=True,  # 以文本模式解码
            check=True,  # 命令失败时抛出异常
        )
        return result.stdout
    except Exception as e:
        print(f"  .doc解析失败（本机缺少textutil或LibreOffice）: {e}")
        return ""


def load_docx(file_path):
    """解析Word(.docx)文件：python-docx 逐段落提取文本"""
    document = docx.Document(file_path)
    text_parts = []
    for paragraph in document.paragraphs:
        text_parts.append(paragraph.text)
    return "\n".join(text_parts)


def load_pptx(file_path):
    """
    解析PPT(.pptx)文件：python-pptx 逐幻灯片、逐形状提取文本
    形状(shape)可能是文本框/图片/表格等，只处理含文本的部分
    """
    presentation = pptx.Presentation(file_path)
    text_parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_parts.append(shape.text_frame.text)
    return "\n".join(text_parts)


def load_xlsx(file_path):
    """
    解析Excel(.xlsx)文件：openpyxl 逐工作表、逐行读取
    每行单元格用" | "分隔，保持表格的行列结构信息
    （data_only=True 表示读取公式的计算结果而不是公式本身）
    """
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    text_parts = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            # 空单元格显示为None，统一转成空字符串再拼接
            row_values = [str(cell) if cell is not None else "" for cell in row]
            text_parts.append(" | ".join(row_values))
    workbook.close()
    return "\n".join(text_parts)


def load_csv(file_path):
    """解析CSV文件：Python内置csv模块逐行读取，单元格用" | "分隔"""
    text_parts = []
    # newline="" 是csv模块官方推荐的打开方式，避免空行问题
    with open(file_path, "r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            text_parts.append(" | ".join(row))
    return "\n".join(text_parts)


def load_html(file_path):
    """解析HTML文件：BeautifulSoup 去掉所有标签，提取纯文本"""
    with open(file_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    # separator="\n" 让原本分块显示的文本之间用换行分隔，避免挤成一行
    return soup.get_text(separator="\n")


def load_xml(file_path):
    """解析XML文件：内置 ElementTree 提取所有文本节点"""
    tree = ET.parse(file_path)
    root = tree.getroot()
    text_parts = []
    # itertext() 深度优先遍历所有文本内容
    for text in root.itertext():
        text = text.strip()
        # 跳过纯空白的文本片段
        if text:
            text_parts.append(text)
    return "\n".join(text_parts)


# ======================================================================
# 第二部分：格式分发表 + 统一入口
# 对应课程代码的 DOCUMENT_LOADER_MAPPING：课程是"扩展名→(Loader类,参数)"，
# 我们是"扩展名→解析函数"，思想完全一致
# ======================================================================

DOCUMENT_PARSER_MAPPING = {
    ".pdf": load_pdf,
    ".txt": load_txt,
    ".doc": load_doc,
    ".docx": load_docx,
    # 注意：旧版二进制.ppt格式需要LibreOffice转换，本课程暂不支持（测试数据中也无此格式）
    ".pptx": load_pptx,
    ".xlsx": load_xlsx,
    ".csv": load_csv,
    ".md": load_md,
    ".xml": load_xml,
    ".html": load_html,
}


def load_document(file_path):
    """
    解析各种文档格式的文件，返回文档内容字符串
    :param file_path: 文档文件路径
    :return: 文档内容的字符串（解析失败或格式不支持时返回空字符串）
    """
    # 获取文件扩展名并转小写，确定文档类型（兼容 .PDF / .Pdf 等写法）
    ext = os.path.splitext(file_path)[1].lower()
    # 从分发表中获取对应的解析函数
    parser_function = DOCUMENT_PARSER_MAPPING.get(ext)

    if parser_function:
        # 调用解析函数得到文档内容字符串
        content = parser_function(file_path)
        # 打印前100个字符，观察各格式的解析效果
        content_preview = content[:100].replace("\n", " ")
        print(f"  解析成功，内容预览: {content_preview}...")
        return content

    print(f"  不支持的文档类型: {ext}")
    return ""


# ======================================================================
# 第三部分：RAG三流程（索引流程升级为文件夹遍历，检索/生成与第2课一致）
# ======================================================================


def load_embedding_model():
    """
    加载本地 bge-small-zh-v1.5 模型（512维中文Embedding模型，离线运行）
    :return: 返回加载好的 SentenceTransformer 模型对象
    """
    print("加载Embedding模型中...")
    embedding_model = SentenceTransformer(str(EMBEDDING_MODEL_PATH))
    print(f"模型最大输入长度(max_seq_length): {embedding_model.max_seq_length}")
    return embedding_model


def indexing_process(folder_path, embedding_model):
    """
    索引流程（第3课升级版）：遍历文件夹中所有文档文件，解析→分块→汇总→向量化→存入FAISS索引。

    :param folder_path: 文档文件夹路径
    :param embedding_model: 预加载的Embedding模型
    :return: 返回 (FAISS向量索引, 所有文档的文本块总列表)
    """
    # 初始化总chunks列表，用于存储所有文档文件的文本块
    all_chunks = []

    # 遍历文件夹中的所有文件。sorted()让遍历顺序稳定，每次运行结果可复现
    for filename in sorted(os.listdir(folder_path)):
        file_path = os.path.join(folder_path, filename)
        # 跳过子目录等非文件条目
        if not os.path.isfile(file_path):
            continue

        print(f"\n处理文档: {filename}")
        # 第1步：按扩展名选择解析器，获得文档的字符串内容
        document_text = load_document(file_path)
        # 解析失败的文件（返回空内容）直接跳过
        if not document_text.strip():
            print("  内容为空，跳过")
            continue
        print(f"  总字符数: {len(document_text)}")

        # 第2步：分块（与第2课相同的递归字符分割器）
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
        )
        chunks = text_splitter.split_text(document_text)
        print(f"  分割的文本Chunk数量: {len(chunks)}")

        # 第3步：将该文档的chunks汇总到总列表
        all_chunks.extend(chunks)

    print(f"\n全部文档处理完毕，共 {len(all_chunks)} 个文本块，开始向量化")

    # 第4步：向量化（与第2课一致：归一化后内积等于余弦相似度）
    embeddings = []
    for chunk in all_chunks:
        embedding = embedding_model.encode(chunk, normalize_embeddings=True)
        embeddings.append(embedding)

    # 转为float32的numpy数组，FAISS要求numpy数组输入
    embeddings_np = np.array(embeddings, dtype="float32")

    # 第5步：构建FAISS索引并加入所有向量
    dimension = embeddings_np.shape[1]
    index = faiss.IndexFlatIP(dimension)
    index.add(embeddings_np)
    print("索引流程完成\n")

    return index, all_chunks


def retrieval_process(query, index, chunks, embedding_model, top_k=TOP_K):
    """
    检索流程：将用户查询转化为向量，在FAISS索引中检索最相似的前K个文本块
    （与第2课完全一致）
    """
    # 查询向量化：必须与索引阶段用同一个Embedding模型（同一向量空间）
    query_embedding = embedding_model.encode(query, normalize_embeddings=True)
    query_embedding = np.array([query_embedding])

    # FAISS检索Top-K：返回相似度得分和命中文本块下标
    scores, indices = index.search(query_embedding, top_k)

    print(f"查询语句: {query}")
    print(f"最相似的前{top_k}个文本块:")

    results = []
    for i in range(top_k):
        result_chunk = chunks[indices[0][i]]
        result_score = scores[0][i]
        print(f"\n--- 文本块{i} (相似度 {result_score:.4f}) ---\n{result_chunk}")
        results.append(result_chunk)

    print("\n检索流程完成")
    return results


def generate_process(query, chunks):
    """
    生成流程：把检索到的文本块与用户问题组装成Prompt，调用DeepSeek生成回答
    （与第2课完全一致）
    """
    # 构建参考文档内容，给每块编号方便模型引用
    context = ""
    for i, chunk in enumerate(chunks):
        context += f"【参考文档{i + 1}】\n{chunk}\n\n"

    prompt = f"请根据以下参考文档回答问题。\n\n{context}\n问题：{query}"
    print(f"生成模型的Prompt:\n{'-' * 60}\n{prompt}\n{'-' * 60}")

    # DeepSeek的OpenAI兼容端点，无需厂商专有SDK
    llm = ChatOpenAI(
        model=os.environ.get("DEEPSEEK_MODEL", "deepseek-flash"),
        api_key=os.environ["DEEPSEEK_API_KEY"],
        base_url="https://api.deepseek.com",
        temperature=0.3,
    )

    messages = [
        (
            "system",
            "你是一个严谨的知识库问答助手。只依据用户提供的参考文档回答；"
            "若参考文档不足以回答问题，请直接说明，不要编造。",
        ),
        ("human", prompt),
    ]

    print("\n生成过程开始:")
    generated_response = ""
    try:
        # 流式输出：每收到一小段就立即打印
        for piece in llm.stream(messages):
            content = piece.content
            generated_response += content
            print(content, end="", flush=True)
        print("\n生成流程完成")
        return generated_response
    except Exception as e:
        print(f"\n大模型调用失败: {e}")
        raise


def main():
    # 检查API Key是否已配置
    if not os.environ.get("DEEPSEEK_API_KEY"):
        raise SystemExit(
            "未检测到 DEEPSEEK_API_KEY。请编辑 .env 文件填入你的Key，"
            "然后用 `uv run --env-file .env python lesson3_multiformat_rag.py` 运行"
        )

    print("RAG流程开始\n" + "=" * 60)
    # 10种格式的测试文档内容完全相同，用于观察不同格式的解析效果
    query = "下面报告中涉及了哪几个行业的案例以及总结各自面临的挑战？"

    embedding_model = load_embedding_model()

    # 索引流程：遍历data/lesson3文件夹中所有格式的文档
    index, chunks = indexing_process(str(DATA_FOLDER), embedding_model)

    retrieval_chunks = retrieval_process(query, index, chunks, embedding_model)
    generate_process(query, retrieval_chunks)

    print("\n" + "=" * 60 + "\nRAG流程结束")


if __name__ == "__main__":
    main()
